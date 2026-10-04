# Importa la hoja "Prácticas" de un Excel de asistencia a la base LOCAL.
# El Excel y los datos NO van al repositorio (es público): se guardan en supabase/privado/ (ignorada por Git).
#   python scripts/importar_asistencia.py            (carga solo si la base no tiene prácticas)
#   python scripts/importar_asistencia.py --reemplazar   (borra las prácticas de la base y las vuelve a cargar)
#   python scripts/importar_asistencia.py --solo-lugares (no importa nada: completa el lugar de las prácticas que no lo tienen)
# El lugar sale del día de la semana, con los mismos horarios de entrenamiento del sitio (TEAM.entrenamientos en js/data.js).
# Reglas del Excel: 1 = fue, 0 = faltó, celda vacía = esa práctica no contaba para la persona (aún no estaba).
import json, os, subprocess, sys, datetime, urllib.request, urllib.error
import openpyxl

RAIZ = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
ARCHIVO = os.path.join(RAIZ, "supabase", "privado", "Asistencia a practicas.xlsx")
HOJA = "Prácticas"

# Nombre en el Excel -> jugador de la base (slug). Revisar si cambia el Excel.
NOMBRES = {
    "santi": "santiago-rodriguez", "mati": "matilde-rodriguez", "juano": "juanjo-alonso", "cami": "camila-couture",
    "thiago": "thiago-elizalde", "seba": "sebastian-migdal", "lea": "leandro-rodriguez", "nico": "nicolas-cabana",
    "sofi": "sofia-rodriguez", "julia": "julieta-noguez", "aini": "ainara-rodriguez", "rosi": "rosina-cordero",
    "luisito": "luis-davila",
}

def entorno():
    salida = subprocess.run("npx --yes supabase status -o env", cwd=RAIZ, shell=True, capture_output=True, text=True, encoding="utf-8").stdout
    env = {}
    for linea in salida.splitlines():
        if "=" in linea:
            k, v = linea.split("=", 1); env[k.strip()] = v.strip().strip('"')
    return env

def pedir(metodo, ruta, cuerpo=None, extra=None):
    cab = {"apikey": KEY, "Authorization": f"Bearer {KEY}", "Content-Type": "application/json", **(extra or {})}
    req = urllib.request.Request(f"{URL}/rest/v1/{ruta}", method=metodo, headers=cab, data=json.dumps(cuerpo).encode() if cuerpo is not None else None)
    try:
        with urllib.request.urlopen(req) as r:
            t = r.read().decode(); return json.loads(t) if t else None
    except urllib.error.HTTPError as e:
        print("Error", e.code, e.read().decode()); sys.exit(1)

env = entorno()
URL, KEY = env.get("API_URL"), env.get("SERVICE_ROLE_KEY")
if not URL or not KEY: sys.exit("No pude leer la base local. ¿Está prendida? (npx supabase start)")
if not URL.startswith(("http://127.0.0.1", "http://localhost")): sys.exit("Esto solo corre contra la base local.")
if not os.path.exists(ARCHIVO): sys.exit(f"No encuentro el Excel en {ARCHIVO}")

DIAS_SEMANA = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]

def completar_lugares():
    """Pone el lugar a las prácticas que no lo tienen, según el día de la semana."""
    salida = subprocess.run(["node", "-e", "const vm=require('vm'),fs=require('fs');const {TEAM}=vm.runInNewContext(fs.readFileSync('js/data.js','utf8')+';({TEAM})');console.log(JSON.stringify(TEAM.entrenamientos))"],
                            cwd=RAIZ, capture_output=True, text=True, encoding="utf-8").stdout
    lugar_por_dia = {e["dia"]: e["lugar"] for e in json.loads(salida)}
    sin_lugar = pedir("GET", "practicas?select=id,fecha&lugar=is.null")
    grupos, sin_asignar = {}, []
    for p in sin_lugar:
        dia = DIAS_SEMANA[datetime.date.fromisoformat(p["fecha"]).weekday()]
        if dia in lugar_por_dia: grupos.setdefault(lugar_por_dia[dia], []).append(p["id"])
        else: sin_asignar.append(p["fecha"])
    for lugar, ids in grupos.items():
        pedir("PATCH", f"practicas?id=in.({','.join(map(str, ids))})", {"lugar": lugar})
        print(f"  {len(ids):2} prácticas → {lugar}")
    if sin_asignar: print("  Sin lugar fijo (se completan a mano en el panel):", ", ".join(sin_asignar))

if "--solo-lugares" in sys.argv:
    completar_lugares(); sys.exit(0)

ws = openpyxl.load_workbook(ARCHIVO, data_only=True)[HOJA]
fechas = {i: c.value.date() for i, c in enumerate(ws[2]) if isinstance(c.value, datetime.datetime)}
filas, totales_hoja = [], {}
for fila in ws.iter_rows(min_row=3, values_only=True):
    nombre = (fila[0] or "").strip()
    if not nombre: continue
    if nombre.lower().startswith("cantidad"):
        totales_hoja = {i: fila[i] for i in fechas}; continue
    filas.append((nombre, fila))

slugs = {j["slug"] for j in pedir("GET", "jugadores?select=slug")}
registros, sin_mapear = {}, []     # fecha -> [(slug, presente)]
for nombre, fila in filas:
    slug = NOMBRES.get(nombre.lower())
    if slug not in slugs: sin_mapear.append(nombre); continue
    for i, f in fechas.items():
        v = fila[i]
        if v in (None, ""): continue                     # no contaba
        registros.setdefault(f, []).append((slug, bool(float(v))))
if sin_mapear: sys.exit(f"Estos nombres del Excel no tienen jugador en la base: {sin_mapear}")

# Control: la fila "Cantidad total" de la hoja debe coincidir con lo que se va a cargar
for i, f in fechas.items():
    esperado, calculado = totales_hoja.get(i), sum(1 for _, p in registros.get(f, []) if p)
    if esperado is not None and int(esperado) != calculado:
        print(f"Aviso: {f} la hoja dice {esperado} asistentes y yo cuento {calculado}")

existentes = pedir("GET", "practicas?select=id&limit=1")
if existentes and "--reemplazar" not in sys.argv: sys.exit("La base ya tiene prácticas. Usá --reemplazar para borrarlas y volver a cargar.")
if existentes: pedir("DELETE", "practicas?id=gt.0")

creadas = pedir("POST", "practicas", [{"fecha": f.isoformat()} for f in sorted(fechas.values())], {"Prefer": "return=representation"})
ids = {p["fecha"]: p["id"] for p in creadas}
filas_asis = [{"practica_id": ids[f.isoformat()], "jugador_slug": s, "presente": p} for f, lista in registros.items() for s, p in lista]
for k in range(0, len(filas_asis), 500): pedir("POST", "asistencias", filas_asis[k:k + 500])
print(f"✓ {len(creadas)} prácticas y {len(filas_asis)} asistencias cargadas")
completar_lugares()
print("Mapeo de nombres:", ", ".join(f"{n.strip()} → {NOMBRES[n.strip().lower()]}" for n, _ in filas))
